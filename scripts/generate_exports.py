from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
import os

data = [
    {
        "title": "1. Project Name & Goal",
        "name": "Project Name: StudyBuddy Learning Platform",
        "goal": "Goal: To create a secure, feature-rich online learning environment combining rich-text note-taking, real-time community interaction, and deeply secure assessment systems augmented by artificial intelligence."
    },
    {
        "title": "2. What Problem Does It Solve",
        "content": "Traditional online learning often requires fragmented tools (one app for notes, one for chat, one for taking tests). StudyBuddy consolidates these functionalities into a single coherent ecosystem, solving the context-switching fatigue for students and enhancing exam security for educators."
    },
    {
        "title": "3. Your Role",
        "content": "Software Engineer / Full Stack Developer\n\nActed as the sole architect for both the frontend (React/Vite) and backend schema (Supabase), managing database modeling, UI/UX design, real-time socket connections, and integration of the Gemini AI API."
    },
    {
        "title": "4. Tech Stack",
        "content": "• Frontend: React 18, Vite, React Router, Tailwind CSS 4\n• Backend & Database: Supabase (PostgreSQL, Realtime WebSockets, Authentication)\n• Rich Text & Media: TipTap Editor, Recharts, html2pdf.js\n• Artificial Intelligence: Google Gemini GenAI SDK"
    },
    {
        "title": "5. Key Features",
        "content": "• Secure Testing Engine: Fullscreen API lockdown, Tab-switching detection, and Server-synced countdown timers.\n• TipTap 'MyNotes': A highly capable rich-text processing suite supporting tables, highlights, and direct PDF exports.\n• Community Hub: Real-time bi-directional study groups and direct messaging powered by WebSockets.\n• Study Assistant AI: Context-aware AI bot capable of formatted markdown streaming."
    },
    {
        "title": "6. Challenge + Solution",
        "content": "Challenge 1: Tracking student cheating and ensuring testing integrity strictly through a web browser.\nSolution: Implemented native browser APIs (visibilitychange and Fullscreen). The wrapper actively listens for backgrounding events; accumulating 3 warnings forces an automatic secure test submission.\n\nChallenge 2: State lag during massive assessment creations.\nSolution: Converted complex questionnaire builder forms from standard useState to a deep useReducer pattern, preventing React from dropping frames during huge re-renders."
    },
    {
        "title": "7. Results / Learning",
        "content": "• Successfully delivered a hardened, complex real-time application with zero active compiler errors.\n• Mastered advanced PostgreSQL Row-Level Security (RLS) policies for scaling multi-tenant data safely.\n• Gained deep, practical experience implementing headless UI architectures and streaming response protocols."
    }
]

def create_docx(filepath):
    doc = Document()
    doc.add_heading('Project Explanation: StudyBuddy', 0)
    
    for item in data:
        doc.add_heading(item['title'], level=1)
        if "name" in item:
            doc.add_paragraph(item['name'])
            doc.add_paragraph(item['goal'])
        else:
            doc.add_paragraph(item['content'])
            
    doc.save(filepath)

def create_pptx(filepath):
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "StudyBuddy Learning Platform"
    subtitle.text = "Project Explanation & Review"
    
    # Content Slides
    bullet_slide_layout = prs.slide_layouts[1]
    
    for item in data:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = item['title']
        tf = body_shape.text_frame
        
        if "name" in item:
            tf.text = item['name']
            p = tf.add_paragraph()
            p.text = item['goal']
        else:
            # Handle bullet points properly
            lines = item['content'].split('\n')
            tf.text = lines[0]
            for line in lines[1:]:
                if line.strip():
                    p = tf.add_paragraph()
                    p.text = line.strip()

    prs.save(filepath)

if __name__ == '__main__':
    docs_dir = r"c:\.webvaradh\IDEA2\studybuddy-learning-platform\docs"
    os.makedirs(docs_dir, exist_ok=True)
    
    docx_path = os.path.join(docs_dir, "StudyBuddy_Project_Explanation.docx")
    pptx_path = os.path.join(docs_dir, "StudyBuddy_Project_Presentation.pptx")
    
    create_docx(docx_path)
    create_pptx(pptx_path)
    
    print(f"Generated: {docx_path}")
    print(f"Generated: {pptx_path}")
